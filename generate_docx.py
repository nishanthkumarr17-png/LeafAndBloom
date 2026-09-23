from docx import Document
from pathlib import Path
import re

md_path = Path('LeafAndBloom_ProjectReport.md')
docx_path = Path('LeafAndBloom_ProjectReport.docx')

text = md_path.read_text(encoding='utf-8')
lines = text.splitlines()
doc = Document()

for line in lines:
    if line.startswith('# '):
        doc.add_heading(line[2:].strip(), level=1)
    elif line.startswith('## '):
        doc.add_heading(line[3:].strip(), level=2)
    elif line.startswith('### '):
        doc.add_heading(line[4:].strip(), level=3)
    elif re.match(r'^[-*] ', line):
        doc.add_paragraph(line[2:].strip(), style='List Paragraph')
    elif line.strip() == '':
        doc.add_paragraph('')
    else:
        doc.add_paragraph(line)

doc.save(docx_path)
print('Created', docx_path.resolve())
